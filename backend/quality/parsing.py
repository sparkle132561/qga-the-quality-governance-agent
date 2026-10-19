import csv
import io
import re
import threading
import zipfile
from functools import lru_cache
from html.parser import HTMLParser

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from PIL import Image
from pypdf import PdfReader

from quality.cleaning import clean_units, normalize_quantity
from quality.pdf_ocr import render_pdf_page

MAX_UNITS = 2000
MAX_TEXT = 250000
SUPPORTED = {".pdf", ".xlsx", ".csv", ".html", ".htm", ".png", ".jpg", ".jpeg"}
_ocr_lock = threading.Lock()


class TextHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.fragments = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.hidden += 1

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, text):
        if not self.hidden and text.strip():
            self.fragments.append(text.strip())


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("文本编码无法识别，请使用UTF-8或GB18030")


@lru_cache(maxsize=1)
def ocr_engine():
    from rapidocr_onnxruntime import RapidOCR

    return RapidOCR(intra_op_num_threads=2, inter_op_num_threads=2)


def parse_material(data: bytes, suffix: str) -> dict:
    if suffix not in SUPPORTED:
        raise ValueError("不支持的文件类型")
    units = []
    total_chars = 0

    def add(text: str, location: dict, confidence: float | None = None):
        nonlocal total_chars
        text = text.strip()
        if not text:
            return
        total_chars += len(text)
        if len(units) >= MAX_UNITS or total_chars > MAX_TEXT:
            raise ValueError("解析内容超出首版处理上限，请拆分文件")
        units.append({"text": text, "location": location, "confidence": confidence})

    engine = "native"
    pdf_pages = []
    page_issues = []
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted or len(reader.pages) > 80:
            raise ValueError("暂不支持加密PDF或超过80页的文件")
        for page_number, page in enumerate(reader.pages, start=1):
            native = page.extract_text() or ""
            # 图片页整页OCR，不能因一个电子页眉而跳过下面的扫描正文。
            if native.strip() and not len(page.images):
                add(native, {"kind": "pdf", "page": page_number, "method": "native"})
                pdf_pages.append({"page": page_number, "method": "native", "status": "parsed"})
                continue
            engine = "pypdfium2-4.30.0+rapidocr-onnxruntime-1.4.4"
            before = len(units)
            try:
                rendered, metadata = render_pdf_page(data, page_number)
                with _ocr_lock:
                    lines, _ = ocr_engine()(rendered)
            except Exception as error:
                # 已成功的页保留；错误页明确待复核，不把部分成功称为完整解析。
                page_issues.append(f"第{page_number}页OCR失败（{type(error).__name__}）")
                add(native, {"kind": "pdf", "page": page_number, "method": "native_fallback"})
                pdf_pages.append({"page": page_number, "method": "ocr", "status": "failed"})
                continue
            for box, text, score in lines or []:
                add(text, {"kind": "pdf", "page": page_number, "method": "ocr",
                           "box": [[float(x), float(y)] for x, y in box],
                           "coordinate_system": "rendered_pixels", **metadata}, float(score))
            recognized = len(units) > before
            if not recognized:
                page_issues.append(f"第{page_number}页未识别到文字，请人工确认是否空白或不可读")
            pdf_pages.append({"page": page_number, "method": "ocr", **metadata,
                              "status": "parsed" if recognized else "needs_info",
                              "native_text": native})
    elif suffix == ".xlsx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError("表格解压内容过大")
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for sheet in workbook:
                if sheet.max_row > 5000 or sheet.max_column > 100:
                    raise ValueError("表格行列超过处理上限")
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            add(
                                str(cell.value),
                                {"kind": "sheet", "sheet": sheet.title, "cell": cell.coordinate},
                            )
        finally:
            workbook.close()
    elif suffix == ".csv":
        for row_number, row in enumerate(csv.reader(io.StringIO(decode_text(data))), start=1):
            if row_number > 5000 or len(row) > 100:
                raise ValueError("CSV行列超过处理上限")
            for column, value in enumerate(row, start=1):
                add(value, {"kind": "csv", "row": row_number, "column": column})
    elif suffix in (".html", ".htm"):
        parser = TextHTML()
        parser.feed(decode_text(data))
        for number, text in enumerate(parser.fragments, start=1):
            add(text, {"kind": "html", "fragment": number})
    else:
        with Image.open(io.BytesIO(data)) as image:
            if image.width * image.height > 16_000_000:
                raise ValueError("图片像素超过1600万")
            image.verify()
        with _ocr_lock:
            lines, _ = ocr_engine()(data)
        engine = "rapidocr-onnxruntime-1.4.4"
        for box, text, score in lines or []:
            add(
                text, {"kind": "image", "box": [[float(x), float(y)] for x, y in box]}, float(score)
            )

    fields = []
    patterns = {
        "model": r"(?:型号|MODEL)\s*[:：]?\s*([A-Za-z0-9][A-Za-z0-9._-]*)",
        "power": r"(?:功率|POWER)\s*[:：]?\s*(\d+(?:\.\d+)?\s*[Ww])",
        "certificate_number": (
            r"(?:证书编号|CERTIFICATE(?:\s+NUMBER|\s+NO\.?|_NUMBER))"
            r"\s*[:：]?\s*([A-Za-z0-9][A-Za-z0-9._-]*)"
        ),
        "net_quantity": (
            r"(?:净含量|净重|NET\s*(?:WEIGHT|WT\.?|QUANTITY)|POIDS\s+NET)"
            r"\s*[:：]?\s*((?:\d+\s*[x×*]\s*)?\d+(?:[.,]\d+)?\s*"
            r"(?:kg|mg|g|ml|cl|l|pieces|piece|pcs|pc))\b"
        ),
    }
    for number, unit in enumerate(units):
        for name, pattern in patterns.items():
            for match in re.finditer(pattern, unit["text"], re.IGNORECASE):
                fields.append(
                    {
                        "name": name,
                        "value": (match.group(1) if name == "net_quantity"
                                  else match.group(1).replace(" ", "")),
                        "unit_index": number,
                        "location": {
                            **unit["location"],
                            "char_start": match.start(1),
                            "char_end": match.end(1),
                        },
                        "confidence": unit["confidence"],
                    }
                )
    labels = {
        "model": "model",
        "型号": "model",
        "power": "power",
        "功率": "power",
        "manufacturer": "manufacturer",
        "生产者": "manufacturer",
        "厂家": "manufacturer",
        "brand": "brand",
        "品牌": "brand",
        "certificate": "certificate_number",
        "证书编号": "certificate_number",
        "生产日期": "production_date",
        "production_date": "production_date",
        "quantity": "net_quantity",
        "net_quantity": "net_quantity",
        "net weight": "net_quantity",
        "净含量": "net_quantity",
        "净重": "net_quantity",
    }
    table_cells = {}
    for index, unit in enumerate(units):
        loc = unit["location"]
        if loc["kind"] == "csv":
            key = ("csv", "", loc["row"], loc["column"])
        elif loc["kind"] == "sheet":
            column, row = coordinate_from_string(loc["cell"])
            key = ("sheet", loc["sheet"], row, column_index_from_string(column))
        else:
            continue
        table_cells[key] = (index, unit)
    header_groups = {}
    uncertain_groups = set()

    def plausible_value(name, text):
        if name in ("model", "certificate_number"):
            return bool(re.fullmatch(r"[A-Za-z0-9._-]+", text) and re.search(r"\d", text))
        if name == "power":
            return bool(re.fullmatch(r"\d+(?:\.\d+)?\s*[Ww]", text))
        if name == "production_date":
            return bool(re.fullmatch(r"\d{4}[-/.]\d{1,2}[-/.]\d{1,2}", text))
        if name == "net_quantity":
            return normalize_quantity(text)["status"] == "normalized"
        return False

    for kind, sheet in {(key[0], key[1]) for key in table_cells}:
        rows = {}
        for (k, s, row, column), (_, unit) in table_cells.items():
            if (k, s) == (kind, sheet):
                rows.setdefault(row, {})[column] = unit["text"].strip()
        first_number = min(rows)
        first_row = rows[first_number]
        known = sum(text.lower() in labels for text in first_row.values())
        later = [values for row, values in rows.items() if row > first_number]
        later_keys = any(values[min(values)].lower() in labels for values in later)
        if later_keys:
            if known >= 2:
                uncertain_groups.add((kind, sheet))
            continue
        if known >= 2:
            header_groups[(kind, sheet)] = first_number
            continue
        column = min(first_row)
        name = labels.get(first_row[column].lower())
        right = first_row.get(column + 1, "")
        if name and plausible_value(name, right):
            continue
        if (
            name
            and later
            and all(column in values and plausible_value(name, values[column]) for values in later)
        ):
            header_groups[(kind, sheet)] = first_number
        elif name and not later and right and right.lower() not in labels:
            continue
        else:
            uncertain_groups.add((kind, sheet))
    for (kind, sheet, row, column), (_, unit) in table_cells.items():
        if (kind, sheet) in uncertain_groups:
            continue
        if (kind, sheet) in header_groups and row != header_groups[(kind, sheet)]:
            continue
        name = labels.get(unit["text"].strip().lower())
        if not name:
            continue
        if (kind, sheet) in header_groups:
            candidate_keys = [
                key
                for key in table_cells
                if key[0:2] == (kind, sheet)
                and key[2] > header_groups[(kind, sheet)]
                and key[3] == column
            ]
        else:
            candidate_keys = [(kind, sheet, row, column + 1)]
        for key in candidate_keys:
            value_cell = table_cells.get(key)
            if value_cell is None:
                continue
            index, other = value_cell
            if other["text"].strip().lower() in labels:
                continue
            fields.append(
                {
                    "name": name,
                    "value": other["text"],
                    "unit_index": index,
                    "location": other["location"],
                    "label_location": unit["location"],
                    "confidence": None,
                }
            )
    for field in fields:
        if field["name"] == "net_quantity":
            field["normalization"] = normalize_quantity(field["value"])
    return {
        "parser_version": "2",
        "engine": engine,
        "units": units,
        "fields": fields,
        "cleaning": clean_units(units),
        "status": "parsed" if units and not uncertain_groups and not page_issues else "needs_info",
        "pdf_pages": pdf_pages,
        "warning": (
            "；".join(page_issues)
            if page_issues
            else "表格布局不确定，保留原始单元格，请人工确认列标题或键值对应"
            if uncertain_groups
            else None
            if units
            else "没有可提取文本，请补充可读资料"
        ),
    }
