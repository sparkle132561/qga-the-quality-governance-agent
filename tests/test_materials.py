import io

from conftest import headers
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from quality.parsing import parse_material


def test_csv_and_sheet_fields_retain_exact_value_location():
    parsed = parse_material("型号,功率\nK170,1800W\n".encode("utf-8-sig"), ".csv")
    model = next(field for field in parsed["fields"] if field["name"] == "model")
    assert model["value"] == "K170"
    assert model["location"] == {"kind": "csv", "row": 2, "column": 1}
    book = Workbook()
    book.active.title = "检测信息"
    book.active.append(["型号", "K170"])
    output = io.BytesIO()
    book.save(output)
    result = parse_material(output.getvalue(), ".xlsx")
    assert result["fields"][0]["location"]["cell"] == "B1"
    assert result["fields"][0]["value"] == "K170"


def test_multiple_text_fields_in_one_unit_keep_values_and_offsets():
    text = (
        "CERTIFICATE NUMBER: C2025-001 CERTIFICATE NUMBER: C2025-002 "
        "MODEL: K170 MODEL: K180 POWER: 1800W POWER: 1500W"
    )
    parsed = parse_material(("<p>" + text + "</p>").encode(), ".html")
    fields = parsed["fields"]
    assert {field["value"] for field in fields if field["name"] == "certificate_number"} == {
        "C2025-001",
        "C2025-002",
    }
    assert {field["value"] for field in fields if field["name"] == "model"} == {"K170", "K180"}
    assert {field["value"] for field in fields if field["name"] == "power"} == {"1800W", "1500W"}
    for field in fields:
        loc = field["location"]
        assert text[loc["char_start"] : loc["char_end"]] == field["value"]


def test_pdf_page_and_html_hidden_content():
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {
                    NameObject("/F1"): DictionaryObject(
                        {
                            NameObject("/Type"): NameObject("/Font"),
                            NameObject("/Subtype"): NameObject("/Type1"),
                            NameObject("/BaseFont"): NameObject("/Helvetica"),
                        }
                    )
                }
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 18 Tf 40 700 Td (MODEL: K170) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    parsed = parse_material(output.getvalue(), ".pdf")
    assert parsed["fields"][0]["value"] == "K170"
    assert parsed["fields"][0]["location"]["page"] == 1
    html = parse_material(b"<script>ignore</script><p>MODEL: K170</p>", ".html")
    assert all("ignore" not in unit["text"] for unit in html["units"])


def test_header_unknown_column_and_all_data_rows():
    content = b"model,remarks,power\nK170,first,1800W\nK180,second,1500W\n"
    csv = parse_material(content, ".csv")
    assert [(item["name"], item["value"]) for item in csv["fields"]] == [
        ("model", "K170"),
        ("model", "K180"),
        ("power", "1800W"),
        ("power", "1500W"),
    ]
    assert csv["fields"][1]["location"]["row"] == 3
    book = Workbook()
    for row in [
        ["model", "remarks", "power"],
        ["K170", "first", "1800W"],
        ["K180", "second", "1500W"],
    ]:
        book.active.append(row)
    output = io.BytesIO()
    book.save(output)
    sheet = parse_material(output.getvalue(), ".xlsx")
    assert [item["value"] for item in sheet["fields"]] == ["K170", "K180", "1800W", "1500W"]
    assert sheet["fields"][1]["location"]["cell"] == "A3"


def test_key_value_csv_and_single_chinese_header_are_not_confused():
    kv = parse_material(b"model,K170\npower,1800W\n", ".csv")
    assert [(field["name"], field["value"]) for field in kv["fields"]] == [
        ("model", "K170"),
        ("power", "1800W"),
    ]
    book = Workbook()
    for row in [["型号", "备注"], ["K170", "first"], ["K180", "second"]]:
        book.active.append(row)
    output = io.BytesIO()
    book.save(output)
    header = parse_material(output.getvalue(), ".xlsx")
    assert [field["value"] for field in header["fields"]] == ["K170", "K180"]
    uncertain = parse_material("型号,不明内容\n不确定值,另一值\n".encode(), ".csv")
    assert uncertain["status"] == "needs_info"
    assert uncertain["fields"] == []
    assert uncertain["units"]


def test_upload_revision_permission_download_and_failure_status(client):
    owner = headers(client)
    case = client.post(
        "/api/cases", headers=owner, json={"title": "文件资料", "workflow": "qualification"}
    ).json()
    path = f"/api/cases/{case['id']}"
    response = client.post(
        path + "/materials",
        headers=owner,
        files={"file": ("label.csv", b"model,power\nK170,1800W\n", "text/csv")},
    )
    assert response.status_code == 201
    material = response.json()
    assert material["status"] == "parsed"
    assert material["case_revision"] == 2
    assert len(material["sha256"]) == 64
    assert client.get(path, headers=owner).json()["revision"] == 2
    other = headers(client, "merchant-b")
    assert client.get(path + "/materials", headers=other).status_code == 404
    download = path + f"/materials/{material['id']}/download"
    assert client.get(download, headers=other).status_code == 404
    assert client.get(download, headers=owner).content == b"model,power\nK170,1800W\n"
    malformed = client.post(
        path + "/materials",
        headers=owner,
        files={"file": ("bad.pdf", b"invalid pdf", "application/pdf")},
    )
    assert malformed.status_code == 201
    assert malformed.json()["status"] == "failed"
    assert (
        client.post(
            path + "/materials",
            headers=owner,
            files={"file": ("bad.exe", b"x", "application/octet-stream")},
        ).status_code
        == 415
    )
    assert client.post(path + "/submit", headers=owner).status_code == 200
    assert (
        client.post(
            path + "/materials", headers=owner, files={"file": ("late.csv", b"x", "text/csv")}
        ).status_code
        == 409
    )
